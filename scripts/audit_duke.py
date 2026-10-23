"""Audit Duke tables offline; export provisional endpoints, never training splits."""

import argparse
from collections import Counter
import csv
import hashlib
from importlib.metadata import version
import json
import math
from pathlib import Path
import statistics

from openpyxl import load_workbook

from download_duke import TABLES, SAMPLE


# Read headers from Data, not Column Explanations: their layouts differ.
FIELDS = {
    "patient_id": "Patient ID",
    "mri": "Days to MRI (From the Date of Diagnosis)",
    "birth": "Date of Birth (Days)",
    "presentation": "Metastatic at Presentation (Outside of Lymph Nodes)",
    "stage_m": "Staging(Metastasis)#(Mx -replaced by -1)[M]",
    "recurrence": "Recurrence event(s)",
    "local": "Days to local recurrence (from the date of diagnosis)",
    "distant": "Days to distant recurrence(from the date of diagnosis)",
    "death": "Days to death (from the date of diagnosis)",
    "last_dr_free": "Days to last distant recurrence free assemssment(from the date of diagnosis)",
    "surgery": "Days to Surgery (from the date of diagnosis)",
}


def numeric(value):
    return (isinstance(value, (int, float)) and not isinstance(value, bool)
            and math.isfinite(value))


def read_clinical(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Data"]
        rows = sheet.iter_rows(values_only=True)
        next(rows)  # Group labels.
        header = [str(v).strip() if v is not None else "" for v in next(rows)]
        indices = {key: header.index(label) for key, label in FIELDS.items()}
        next(rows)  # Codes and explanations, not a patient.
        records = []
        for row in rows:
            if not any(v is not None for v in row):
                continue
            record = {key: row[index] for key, index in indices.items()}
            if not str(record["patient_id"]).startswith("Breast_MRI_"):
                raise ValueError("Unexpected patient ID in Data sheet")
            records.append(record)
        check_unique([r["patient_id"] for r in records], path.name)
        return records, {s.title: {"rows": s.max_row, "columns": s.max_column}
                         for s in workbook}
    finally:
        workbook.close()


def check_unique(ids, name):
    if len(ids) != len(set(ids)):
        raise ValueError(f"Duplicate patient IDs: {name}; join would multiply rows")


def read_auxiliary(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows)
        if header[0] != "Patient ID":
            raise ValueError(f"Unexpected schema: {path.name}")
        values = {"ids": [], "columns": len(header) - 1,
                  "non_numeric_cells": 0, "incomplete_rows": 0}
        for row in rows:
            if not row[0]:
                continue
            values["ids"].append(row[0])
            bad = sum(not numeric(v) for v in row[1:])
            values["non_numeric_cells"] += bad
            values["incomplete_rows"] += bool(bad)
        check_unique(values["ids"], path.name)
        return values
    finally:
        workbook.close()


def endpoint(record):
    """Return an audit endpoint and flags; missing follow-up is not a negative label.

    Death is coded as competing only when disease-free ascertainment reaches it.
    Otherwise censor at the last disease-specific assessment; flag the later death.
    Numeric recurrence dates with conflicting flags remain visible but quarantined.
    """
    flags = []
    mri, dr, death, last = (record[k] for k in ("mri", "distant", "death", "last_dr_free"))
    result = {"time_from_mri_days": None, "event_type": None, "flags": flags}
    if not numeric(mri):
        flags.append("missing_mri_time")
        return result
    if mri < 0:
        flags.append("mri_before_diagnosis")
    if numeric(record["surgery"]) and record["surgery"] < mri:
        flags.append("surgery_before_mri")
    if numeric(dr):
        if record["recurrence"] != 1:
            flags.append("distant_date_conflicts_with_recurrence_flag")
        if numeric(death) and death < dr:
            flags.append("death_before_distant_recurrence")
        end, event_type = dr, 1
    else:
        if record["recurrence"] == 1 and not numeric(record["local"]):
            flags.append("recurrence_without_local_or_distant_date")
        if record["distant"] != "NP":
            flags.append("unknown_distant_event_status")
        if not numeric(last):
            flags.append("missing_disease_specific_followup")
            return result
        end, event_type = last, 0
        if numeric(death):
            if death <= last:
                end, event_type = death, 2
            else:
                flags.append("death_after_last_dr_free_assessment")
    result.update(time_from_mri_days=end - mri, event_type=event_type)
    if end <= mri:
        flags.append("no_positive_followup_after_mri")
    return result


BLOCKING_FLAGS = {
    "missing_mri_time", "distant_date_conflicts_with_recurrence_flag",
    "death_before_distant_recurrence", "recurrence_without_local_or_distant_date",
    "unknown_distant_event_status", "missing_disease_specific_followup",
    "no_positive_followup_after_mri", "surgery_before_mri",
}


def horizon_counts(rows, years):
    counts = Counter({"distant_recurrence": 0, "competing_death": 0,
                      "observed_event_free": 0, "early_censored": 0})
    horizon = years * 365.25
    for row in rows:
        time, event = row["time_from_mri_days"], row["event_type"]
        if time <= horizon and event == 1:
            counts["distant_recurrence"] += 1
        elif time <= horizon and event == 2:
            counts["competing_death"] += 1
        elif time >= horizon:
            counts["observed_event_free"] += 1
        else:
            counts["early_censored"] += 1
    return dict(counts)


def summarize(rows):
    return {
        "patients": len(rows),
        "distant_events": sum(r["event_type"] == 1 for r in rows),
        "competing_deaths_observed": sum(r["event_type"] == 2 for r in rows),
        "median_endpoint_time_days": statistics.median(
            r["time_from_mri_days"] for r in rows) if rows else None,
        "flags": dict(Counter(f for r in rows for f in r["flags"])),
        "horizons_years": {str(y): horizon_counts(rows, y) for y in (1, 2, 3, 5)},
    }


def write_csv(path, rows, fields):
    with path.open("w", newline="") as output:
        writer = csv.DictWriter(output, fields)
        writer.writeheader()
        writer.writerows({k: ";".join(row[k]) if isinstance(row[k], list) else row[k]
                          for k in fields} for row in rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=Path("data/raw/duke"))
    parser.add_argument("--output-dir", type=Path, default=Path("reports/duke"))
    parser.add_argument("--derived-dir", type=Path, default=Path("data/derived/duke"))
    args = parser.parse_args()
    records, sheets = read_clinical(args.data_dir / "Clinical_and_Other_Features.xlsx")
    clinical_ids = {r["patient_id"] for r in records}
    auxiliary = {}
    for name in ("Annotation_Boxes.xlsx", "Imaging_Features.xlsx"):
        info = read_auxiliary(args.data_dir / name)
        ids = set(info.pop("ids"))
        auxiliary[name] = info | {"patients": len(ids),
            "missing_clinical_ids": len(clinical_ids - ids),
            "extra_ids": len(ids - clinical_ids)}
        if ids != clinical_ids:
            raise ValueError(f"Patient IDs do not match: {name}")

    audited = []
    for record in records:
        derived = endpoint(record)
        time = derived["time_from_mri_days"]
        positive = time is not None and time > 0
        non_m1 = record["presentation"] == 0 and record["stage_m"] in (0, -1)
        strict = non_m1 and record["stage_m"] == 0 and positive
        reviewed = not (set(derived["flags"]) & BLOCKING_FLAGS)
        audited.append(record | derived | {"strict_m0": strict,
            "extended_non_m1": non_m1 and positive,
            "passes_outcome_qc": strict and reviewed})

    strict = [r for r in audited if r["strict_m0"]]
    reviewed = [r for r in audited if r["passes_outcome_qc"]]
    extended = [r for r in audited if r["extended_non_m1"]]
    summary = {
        "status": "provisional_audit_not_training_ready",
        "clinical_patients": len(records), "sheets": sheets,
        "stage_m_counts": dict(Counter(str(r["stage_m"]) for r in records)),
        "metastatic_at_presentation": sum(r["presentation"] == 1 for r in records),
        "stage_m1_presentation_no": sum(r["stage_m"] == 1 and r["presentation"] == 0
                                         for r in records),
        "numeric_distant_dates_all": sum(numeric(r["distant"]) for r in records),
        "auxiliary": auxiliary,
        "strict_m0_provisional": summarize(strict),
        "strict_m0_after_outcome_qc": summarize(reviewed),
        "extended_including_mx_provisional": summarize(extended),
    }
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.derived_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    manifest = {"dataset_doi": "10.7937/TCIA.e3sv-re93", "files": [],
                "software": {p: version(p) for p in ("openpyxl", "pydicom")}}
    for name, url in (TABLES | SAMPLE).items():
        path = args.data_dir / name
        if path.exists():
            with path.open("rb") as source:
                digest = hashlib.file_digest(source, "sha256").hexdigest()
            manifest["files"].append({"name": name, "url": url,
                "bytes": path.stat().st_size,
                "sha256": digest})
    (args.output_dir / "source_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    fields = list(FIELDS) + ["time_from_mri_days", "event_type", "strict_m0",
                            "extended_non_m1", "passes_outcome_qc", "flags"]
    write_csv(args.derived_dir / "patient_audit.csv", audited, fields)
    write_csv(args.derived_dir / "cohort_provisional.csv", strict, fields)
    write_csv(args.derived_dir / "needs_review.csv",
              [r for r in strict if not r["passes_outcome_qc"]], fields)
    write_report(args.output_dir / "audit.md", summary)
    print(json.dumps({k: summary[k] for k in ("clinical_patients", "strict_m0_provisional",
                     "strict_m0_after_outcome_qc")}, indent=2))


def write_report(path, summary):
    strict = summary["strict_m0_provisional"]
    reviewed = summary["strict_m0_after_outcome_qc"]
    extended = summary["extended_including_mx_provisional"]
    lines = ["# Duke: воспроизводимый аудит таблиц", "",
        "Отчёт генерирует `scripts/audit_duke.py`. Это проверка данных, не результат обучения.",
        "Источники и SHA-256: [source_manifest.json](source_manifest.json). "
        "Определения полей: [TCIA](https://wiki.cancerimagingarchive.net/pages/viewpage.action?pageId=70226903).", "",
        "## Отбор", "", "| Этап | Пациенты | Даты отдалённого рецидива |",
        "| --- | ---: | ---: |",
        f"| Исходная клиника | {summary['clinical_patients']} | {summary['numeric_distant_dates_all']} |",
        f"| Не M1, включая Mx, положительное время после МРТ | {extended['patients']} | {extended['distant_events']} |",
        f"| Строгий M0, положительное время после МРТ | {strict['patients']} | {strict['distant_events']} |",
        f"| Строгий M0 после исключения противоречий исхода/операции | {reviewed['patients']} | {reviewed['distant_events']} |", "",
        "Строгий M0: `stage M = 0` и `Metastatic at Presentation = 0`. "
        "Mx (-1) не заменяется на M0. Время события/оценки без рецидива пересчитано от МРТ; "
        "нулевое и отрицательное наблюдение исключено. Отрицательный срок самой МРТ от диагноза не исключён автоматически.", "",
        f"M-стадии: `{summary['stage_m_counts']}`. Метастазы при обращении: "
        f"{summary['metastatic_at_presentation']}; M1 при флаге отсутствия метастазов: "
        f"{summary['stage_m1_presentation_no']}.", "",
        "## Сопоставление таблиц", "",
        "| Файл | Уникальные ID | Столбцы без ID | Нечисловые ячейки признаков | Несовпавшие ID |",
        "| --- | ---: | ---: | ---: | ---: |"]
    for name, info in summary["auxiliary"].items():
        lines.append(f"| {name} | {info['patients']} | {info['columns']} | "
                     f"{info['non_numeric_cells']} | {info['missing_clinical_ids'] + info['extra_ids']} |")
    lines += ["", "Число столбцов не доказывает пригодность всех признаков. "
        "Расположение полей читается по заголовкам листа Data: Column Explanations имеет другую ширину. "
        "Дубликаты ID и несовпадения таблиц вызывают ошибку, а не неявное объединение.", "",
        "## Замечания в предварительной M0-когорте", "",
        "| Код проверки | Пациенты |", "| --- | ---: |"]
    for flag, count in sorted(strict["flags"].items()):
        lines.append(f"| `{flag}` | {count} |")
    lines += ["", "Конфликт даты и флага рецидива не исправляется автоматически: "
        "строки остаются в предварительном аудите, но исключены из столбца после QC. "
        "Их ID сохранены локально в `data/derived/duke/needs_review.csv`.", "",
        "Смерть после последней оценки отсутствия отдалённого рецидива не переносит "
        "цензурирование на дату смерти: безрецидивный статус в этом промежутке не подтверждён. "
        "Временный endpoint цензурируется по последней специализированной оценке. "
        "Смерть кодируется как конкурирующая только при наблюдении отсутствия рецидива до неё. "
        "Это консервативное правило аудита; анализ чувствительности и окончательное правило нужны до обучения.", "",
        "## Покрытие горизонтов после QC", "",
        "| Лет от МРТ | Отдалённый рецидив к сроку | Наблюдаемая конкурирующая смерть | Без события к сроку | Раннее цензурирование |",
        "| --- | ---: | ---: | ---: | ---: |"]
    for year, counts in reviewed["horizons_years"].items():
        lines.append(f"| {year} | {counts['distant_recurrence']} | {counts['competing_death']} | "
                     f"{counts['observed_event_free']} | {counts['early_censored']} |")
    lines += ["", "Это описательные количества, не оценки вероятности. Раннее цензурирование "
        "не является отрицательной меткой. Даты рецидивов с противоречивыми флагами в этой таблице исключены.", "",
        "## Что не завершено", "",
        "Доступность рецепторных результатов на дату МРТ, происхождение TNM/grade, "
        "правило конкурирующей смерти и окончательная аналитическая когорта требуют решения. "
        "Отрицательные сроки МРТ от диагноза требуют отдельной проверки доступности клиники. "
        "Разбиения train/test пока не созданы.", ""]
    path.write_text("\n".join(lines))


if __name__ == "__main__":
    main()
