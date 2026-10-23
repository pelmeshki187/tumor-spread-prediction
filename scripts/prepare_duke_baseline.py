"""Prepare a fixed, allowlisted Duke baseline table and shared patient folds."""

import argparse
from collections import Counter
import csv
import hashlib
import json
import math
from pathlib import Path
import statistics

import numpy as np
from openpyxl import load_workbook
from sklearn.model_selection import StratifiedKFold

from audit_duke import audit_records, check_unique, numeric, read_clinical


VOLUME_COLUMN = "Volume_cu_mm_Tumor"
FEATURES = ("age_mri_years", "log1p_tumor_volume_mm3")
SEED = 20261023
N_FOLDS = 5


def sha256(path):
    with path.open("rb") as source:
        return hashlib.file_digest(source, "sha256").hexdigest()


def volume_value(value):
    if value is None or (isinstance(value, str) and value.strip().lower() in {"", "nan", "na", "nc", "np"}):
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if not numeric(value) or value <= 0:
        raise ValueError(f"Invalid tumor volume: {value!r}")
    return float(value)


def read_volumes(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows)
        if header[0] != "Patient ID" or header.count(VOLUME_COLUMN) != 1:
            raise ValueError("Unexpected imaging feature schema")
        index = header.index(VOLUME_COLUMN)
        pairs = [(r[0], volume_value(r[index])) for r in rows if r[0]]
        check_unique([p[0] for p in pairs], path.name)
        return dict(pairs)
    finally:
        workbook.close()


def assign_folds(rows, seed=SEED, n_folds=N_FOLDS):
    rows = sorted(rows, key=lambda r: r["patient_id"])
    check_unique([r["patient_id"] for r in rows], "cohort")
    labels = [r["event_type"] == 1 for r in rows]
    if len(set(labels)) != 2 or min(Counter(labels).values()) < n_folds:
        raise ValueError("Too few patients/events for the requested stratified folds")
    splitter = StratifiedKFold(n_splits=n_folds, shuffle=True, random_state=seed)
    assignment = {}
    for fold, (_, test) in enumerate(splitter.split(np.zeros(len(rows)), labels)):
        for index in test:
            assignment[rows[index]["patient_id"]] = fold
    return assignment


def build_tables(records, volumes):
    if set(volumes) != {r["patient_id"] for r in records}:
        raise ValueError("Clinical/imaging patient IDs do not match")
    cohort = sorted((r for r in audit_records(records) if r["passes_outcome_qc"]),
                    key=lambda r: r["patient_id"])
    features, outcomes = [], []
    for row in cohort:
        age = (row["mri"] - row["birth"]) / 365.25
        if not numeric(age) or not 18 <= age <= 110:
            raise ValueError("Age outside the predeclared adult sanity range [18, 110]")
        volume = volumes[row["patient_id"]]
        features.append({"patient_id": row["patient_id"], "age_mri_years": age,
            "log1p_tumor_volume_mm3": math.log1p(volume) if volume is not None else None})
        outcomes.append({"patient_id": row["patient_id"],
            "time_days": row["time_from_mri_days"], "event_type": row["event_type"],
            "death_gap": int("death_after_last_dr_free_assessment" in row["flags"]),
            "mri_before_diagnosis": int("mri_before_diagnosis" in row["flags"])})
    return cohort, features, outcomes


def save_csv(path, rows, fields):
    with path.open("w", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=Path("data/raw/duke"))
    parser.add_argument("--output-dir", type=Path, default=Path("data/derived/duke/baseline"))
    parser.add_argument("--report-dir", type=Path, default=Path("reports/duke"))
    args = parser.parse_args()
    source_names = ("Clinical_and_Other_Features.xlsx", "Imaging_Features.xlsx")
    sources = {name: sha256(args.data_dir / name) for name in source_names}
    records, _ = read_clinical(args.data_dir / source_names[0])
    volumes = read_volumes(args.data_dir / source_names[1])
    cohort, features, outcomes = build_tables(records, volumes)
    folds = assign_folds(cohort)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.report_dir.mkdir(parents=True, exist_ok=True)
    save_csv(args.output_dir / "features.csv", features, ["patient_id", *FEATURES])
    save_csv(args.output_dir / "outcomes.csv", outcomes,
             ["patient_id", "time_days", "event_type", "death_gap", "mri_before_diagnosis"])
    save_csv(args.output_dir / "folds.csv",
             [{"patient_id": r["patient_id"], "fold": folds[r["patient_id"]]} for r in cohort],
             ["patient_id", "fold"])
    observed_volumes = [volumes[r["patient_id"]] for r in cohort if volumes[r["patient_id"]] is not None]
    fold_counts = []
    for fold in range(N_FOLDS):
        subset = [r for r in outcomes if folds[r["patient_id"]] == fold]
        fold_counts.append({"fold": fold, "patients": len(subset),
            "distant_events": sum(r["event_type"] == 1 for r in subset),
            "distant_events_before_3y": sum(r["event_type"] == 1 and r["time_days"] < 3*365.25 for r in subset)})
    summary = {"status": "exploratory_cause_specific_ranking_only", "seed": SEED,
        "n_folds": N_FOLDS, "patients": len(cohort),
        "distant_events": sum(r["event_type"] == 1 for r in outcomes),
        "feature_allowlist": list(FEATURES), "volume_source_column": VOLUME_COLUMN,
        "volume_unit": "mm^3 as labelled by publisher; not independently recomputed",
        "volume": {"observed": len(observed_volumes), "missing": len(cohort)-len(observed_volumes),
            "minimum": min(observed_volumes), "median": statistics.median(observed_volumes),
            "maximum": max(observed_volumes)},
        "age_range_years": [min(r["age_mri_years"] for r in features), max(r["age_mri_years"] for r in features)],
        "source_sha256": sources,
        "prepared_sha256": {name: sha256(args.output_dir / name)
                             for name in ("features.csv", "outcomes.csv", "folds.csv")},
        "folds": fold_counts,
        "source_definitions": ["https://www.nature.com/articles/s41416-018-0185-8",
                               "https://pmc.ncbi.nlm.nih.gov/articles/PMC6176866/"],
        "missing_value_policy": "No global imputation; training-fold median after log1p",
        "selection": "strict M0, positive follow-up, no blocking outcome/surgery flags"}
    for path in (args.output_dir / "manifest.json", args.report_dir / "baseline_inputs.json"):
        path.write_text(json.dumps(summary, indent=2, allow_nan=False) + "\n")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
