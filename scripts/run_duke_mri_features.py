"""Evaluate a prespecified nine-feature MRI panel with nested ridge Cox tuning."""

import argparse
from collections import Counter
from importlib.metadata import version
import json
import math
from pathlib import Path
import platform
import time
import warnings

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.exceptions import ConvergenceWarning
from sklearn.impute import SimpleImputer
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from sksurv.linear_model import CoxPHSurvivalAnalysis
from sksurv.util import Surv

from audit_duke import check_unique
from check_duke_stability import METRICS, SEEDS, describe
from prepare_duke_baseline import FEATURES, N_FOLDS, assign_folds, save_csv, sha256, volume_value
from run_duke_baseline import HORIZON_DAYS, aggregate_folds, harrell_metrics, ipcw_metrics, load_prepared, make_model


VOLUME = FEATURES[1]
EXTRA_FEATURES = (
    "Median_solidity_Tumor", "Median_Elongation_Tumor",
    "Entropy_Tumor", "Homogeneity1_Tumor",
    "Max_Enhancement_from_char_curv", "Time_to_Peak_from_char_curv",
    "Uptake_rate_from_char_curv", "Washout_rate_from_char_curv")
PANEL = (VOLUME, *EXTRA_FEATURES)
ALPHAS = (0.1, 1.0, 10.0, 100.0, 1000.0)
INNER_FOLDS = 3
MODEL_COLUMNS = {"volume": [VOLUME], "ridge_volume": [VOLUME], "ridge_mri": list(PANEL)}
LABELS = {"volume": "Cox: объём", "ridge_volume": "Ridge Cox: объём",
          "ridge_mri": "Ridge Cox: 9 признаков МРТ"}
SCENARIOS = ("primary", "exclude_death_gap")


def numeric_feature(value):
    if value is None or (isinstance(value, str) and value.strip().lower() in {"", "nan", "na", "nc", "np"}):
        return np.nan
    if isinstance(value, float) and math.isnan(value):
        return np.nan
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError("Unexpected nonnumeric/nonfinite MRI feature")
    return float(value)


def validate_panel(features):
    if tuple(features.columns) != PANEL:
        raise ValueError("MRI table must contain exactly the prespecified panel in order")
    if features.index.has_duplicates or features.index.isna().any():
        raise ValueError("Missing or duplicate MRI patient IDs")
    if np.isinf(features.to_numpy(dtype=float)).any():
        raise ValueError("Infinite MRI feature")


def read_panel(path, baseline, manifest):
    if sha256(path) != manifest["source_sha256"]["Imaging_Features.xlsx"]:
        raise ValueError("Imaging source differs from the baseline source")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows)
        names = ("Patient ID", "Volume_cu_mm_Tumor", *EXTRA_FEATURES)
        if any(header.count(name) != 1 for name in names):
            raise ValueError("Missing or duplicated MRI source column")
        positions = [header.index(name) for name in names]
        records = []
        for row in rows:
            if row[positions[0]] is None:
                continue
            volume = volume_value(row[positions[1]])
            records.append({"patient_id": row[positions[0]],
                VOLUME: math.log1p(volume) if volume is not None else np.nan,
                **{name: numeric_feature(row[index]) for name, index in zip(EXTRA_FEATURES, positions[2:])}})
    finally:
        workbook.close()
    check_unique([r["patient_id"] for r in records], "MRI features")
    frame = pd.DataFrame(records).set_index("patient_id")
    if not set(baseline.index) <= set(frame.index):
        raise ValueError("Baseline patients missing from MRI source")
    frame = frame.loc[baseline.index, list(PANEL)]
    validate_panel(frame)
    if not np.allclose(frame[VOLUME], baseline[VOLUME], rtol=1e-12, atol=1e-12, equal_nan=True):
        raise ValueError("MRI volume does not reproduce the baseline volume")
    return frame


def ridge_model(alpha):
    return make_pipeline(SimpleImputer(strategy="median", keep_empty_features=True), StandardScaler(),
        CoxPHSurvivalAnalysis(alpha=alpha, ties="efron", n_iter=200, tol=1e-9))


def fit_checked(model, features, outcome):
    with warnings.catch_warnings():
        warnings.simplefilter("error", ConvergenceWarning)
        warnings.simplefilter("error", RuntimeWarning)
        model.fit(features, outcome)
    if not np.isfinite(model.steps[-1][1].coef_).all():
        raise ValueError("Nonfinite Cox coefficients")
    return model


def preprocessing_details(model, columns):
    imputer, scaler, cox = (step[1] for step in model.steps)
    return {"imputer_train_medians": dict(zip(columns, imputer.statistics_.tolist())),
            "scaler_train_means": dict(zip(columns, scaler.mean_.tolist())),
            "scaler_train_scales": dict(zip(columns, scaler.scale_.tolist())),
            "cox_coefficients_standardized": dict(zip(columns, cox.coef_.tolist()))}


def inner_assignment(features, outcome, seed):
    rows = [{"patient_id": p, "event_type": int(e)} for p, e in zip(features.index, outcome["event"])]
    return pd.Series(assign_folds(rows, seed=seed, n_folds=INNER_FOLDS)).loc[features.index]


def tune_ridge(features, outcome, seed, alphas=ALPHAS):
    """Receives outer-training data only; every candidate refits inner preprocessing."""
    if not alphas or len(set(alphas)) != len(alphas) or any(not np.isfinite(a) or a <= 0 for a in alphas):
        raise ValueError("Ridge penalties must be positive, finite and unique")
    folds = inner_assignment(features, outcome, seed)
    candidates = []
    for alpha in alphas:
        metrics, processing = [], []
        for fold in range(INNER_FOLDS):
            train, test = np.asarray(folds != fold), np.asarray(folds == fold)
            model = fit_checked(ridge_model(alpha), features.iloc[train], outcome[train])
            scores = model.predict(features.iloc[test])
            if not np.isfinite(scores).all():
                raise ValueError("Nonfinite inner predictions")
            metrics.append({"fold": fold, "train_patients": int(train.sum()),
                "validation_patients": int(test.sum()),
                "train_events": int(outcome[train]["event"].sum()),
                "validation_events": int(outcome[test]["event"].sum()),
                **harrell_metrics(outcome[test], scores)})
            # These statistics are identical for all alpha values, so store them once.
            if alpha == alphas[0]:
                details = preprocessing_details(model, list(features.columns))
                details.pop("cox_coefficients_standardized")
                processing.append({"fold": fold, **details})
        denominator = sum(m["comparable_pairs"] for m in metrics)
        numerator = sum(m["concordant"] + 0.5 * m["tied_risk"] for m in metrics)
        candidates.append({"alpha": alpha, "inner_cindex": numerator / denominator, "folds": metrics})
        if alpha == alphas[0]:
            inner_processing = processing
    best_score = max(c["inner_cindex"] for c in candidates)
    chosen = max(c["alpha"] for c in candidates if best_score - c["inner_cindex"] <= 1e-12)
    model = fit_checked(ridge_model(chosen), features, outcome)
    return model, {"inner_seed": seed, "inner_n_folds": INNER_FOLDS, "selected_alpha": chosen,
                   "selection_metric": "Harrell C by within-inner-fold pair totals",
                   "tie_break": "largest alpha within 1e-12 of best score",
                   "candidates": candidates, "inner_preprocessing": inner_processing}


def fit_outer_models(features, outcome, inner_seed, alphas=ALPHAS):
    """All tuning and refitting happen before any outer-test data are passed in."""
    validate_panel(features)
    models = {}
    baseline = fit_checked(make_model(), features[[VOLUME]], outcome)
    models["volume"] = (baseline, None)
    for name in ("ridge_volume", "ridge_mri"):
        models[name] = tune_ridge(features[MODEL_COLUMNS[name]], outcome, inner_seed, alphas)
    return models


def evaluate_outer(features, outcomes, folds, selected, name, seed, alphas=ALPHAS):
    scenario = {"name": name, "patients": int(np.sum(selected)),
        "distant_events": int(np.sum(selected & (outcomes["event_type"] == 1))),
        "models": {model: {"features": columns, "folds": []} for model, columns in MODEL_COLUMNS.items()}}
    predictions, inner_rows = [], []
    for fold in range(N_FOLDS):
        train, test = np.asarray(selected & (folds != fold)), np.asarray(selected & (folds == fold))
        y_train = Surv.from_arrays(outcomes.iloc[train]["event_type"] == 1, outcomes.iloc[train]["time_days"])
        y_test = Surv.from_arrays(outcomes.iloc[test]["event_type"] == 1, outcomes.iloc[test]["time_days"])
        if not y_train["event"].any() or not y_test["event"].any():
            raise ValueError("Every outer fold needs distant events")
        inner_seed = seed * 10 + fold
        fitted = fit_outer_models(features.iloc[train], y_train, inner_seed, alphas)
        assignment = inner_assignment(features.iloc[train], y_train, inner_seed)
        inner_rows.extend({"seed": seed, "scenario": name, "outer_fold": fold,
            "patient_id": p, "inner_fold": int(assignment.loc[p])} for p in features.index[train])
        for model_name, (model, tuning) in fitted.items():
            columns = MODEL_COLUMNS[model_name]
            scores = model.predict(features.iloc[test][columns])
            if not np.isfinite(scores).all():
                raise ValueError("Nonfinite outer held-out scores")
            diagnostics = {"fold": fold, "train_patients": int(train.sum()), "test_patients": int(test.sum()),
                "train_events": int(y_train["event"].sum()), "test_events": int(y_test["event"].sum()),
                **preprocessing_details(model, columns), "harrell": harrell_metrics(y_test, scores),
                "ipcw_3y": ipcw_metrics(y_train, y_test, scores), "tuning": tuning}
            scenario["models"][model_name]["folds"].append(diagnostics)
            predictions.extend({"seed": seed, "scenario": name, "model": model_name,
                "patient_id": p, "fold": fold, "risk_score": float(score)}
                for p, score in zip(features.index[test], scores))
    for model in scenario["models"].values():
        model["aggregate"] = aggregate_folds(model["folds"])
    expected = scenario["patients"] * len(MODEL_COLUMNS)
    if len(predictions) != expected or len({(r["patient_id"], r["model"]) for r in predictions}) != expected:
        raise ValueError("Each patient must have exactly one outer prediction per model")
    return scenario, predictions, inner_rows


def summarize(repeats):
    result = {}
    for scenario in SCENARIOS:
        result[scenario] = {}
        for metric in METRICS:
            values = {model: [r["scenarios"][scenario]["models"][model]["aggregate"][metric]
                              for r in repeats] for model in MODEL_COLUMNS}
            result[scenario][metric] = {}
            for model, scores in values.items():
                comparisons = {}
                for reference in ("volume", "ridge_volume"):
                    delta = [a - b if a is not None and b is not None else None
                             for a, b in zip(scores, values[reference])]
                    comparisons[reference] = describe(delta) | {
                        "positive_repeats": sum(v is not None and v > 0 for v in delta)}
                result[scenario][metric][model] = {"scores": describe(scores), "paired_differences": comparisons}
    return result


def run_experiment(features, outcomes, seeds=SEEDS, alphas=ALPHAS):
    validate_panel(features)
    if not features.index.equals(outcomes.index):
        raise ValueError("MRI features and outcomes must have identical ordered patient IDs")
    if not seeds or len(set(seeds)) != len(seeds):
        raise ValueError("Seeds must be nonempty and unique")
    rows = [{"patient_id": p, "event_type": int(outcomes.loc[p, "event_type"])} for p in features.index]
    repeats, predictions, assignments, inner_rows = [], [], [], []
    for seed in seeds:
        folds = pd.Series(assign_folds(rows, seed=seed)).loc[features.index]
        assignments.extend({"seed": seed, "patient_id": p, "fold": int(folds.loc[p])} for p in features.index)
        repeat = {"seed": seed, "scenarios": {}}
        for name, selected in (("primary", np.ones(len(features), dtype=bool)),
                               ("exclude_death_gap", outcomes["death_gap"] == 0)):
            scenario, scores, inner = evaluate_outer(features, outcomes, folds, selected, name, seed, alphas)
            repeat["scenarios"][name] = scenario
            predictions.extend(scores)
            inner_rows.extend(inner)
        repeats.append(repeat)
        print(f"Completed nested MRI comparison {len(repeats)}/{len(seeds)} (seed={seed})", flush=True)
    return repeats, predictions, assignments, inner_rows


def write_report(path, result):
    def fmt(stats):
        if stats["mean"] is None:
            return "недоступно"
        return f"{stats['mean']:.3f} [{stats['minimum']:.3f}; {stats['maximum']:.3f}]"

    first = result["repeats"][0]["scenarios"]
    config = result["configuration"]
    lines = ["# Duke: девять признаков МРТ и вложенная настройка Ridge Cox", "",
        f"{first['primary']['patients']} пациентов / {first['primary']['distant_events']} "
        f"отдалённых рецидивов; {len(config['outer_seeds'])} повторений пятиблочной "
        "валидации по тем же пациентам, что в [предыдущем сравнении](model_comparison.md).", "",
        "## Входы и настройка", "",
        "Набор зафиксирован по происхождению и смысловым группам до сравнения с исходами: "
        "log-объём, solidity, elongation, entropy, homogeneity и четыре характеристики "
        "кривой контрастирования (максимум, время пика, uptake и washout). "
        "[Обоснование и первичные источники](../../research/duke-mri-feature-rationale.md). "
        "Возраст и остальные 520 опубликованных признаков не используются. "
        "Точные единицы/формулы части полей независимо не подтверждены; это анализ "
        "опубликованных числовых дескрипторов исходной МРТ, без пересчёта по DICOM.", "",
        "Три модели: прежняя Cox по объёму без штрафа, Ridge Cox по объёму и Ridge Cox "
        "по всем девяти признакам. Для обеих Ridge-моделей alpha выбирается независимо "
        f"из {config['alphas']} на трёх внутренних стратифицированных блоках "
        "исключительно внешней обучающей выборки. Критерий — Harrell C по сумме пар "
        "внутри внутренних блоков; при равенстве в пределах 1e-12 выбирается наибольшая "
        "alpha. Внутренний seed = 10 × внешний seed + номер внешнего блока.", "",
        "Медиана, стандартизация и модель заново обучаются внутри каждого внутреннего "
        "блока и затем на всей внешней обучающей части с выбранной alpha. Все девять "
        "признаков входят в Ridge; штраф сжимает коэффициенты, автоматического отбора "
        "по исходам нет. Постоянные столбцы сохраняются; полностью пустой обучающий "
        "столбец Ridge заполняется нулём. Cox: ties=Efron, n_iter=200 для Ridge, "
        "tol=1e-9. Ошибки/предупреждения сходимости прерывают расчёт и не заменяются "
        "числом качества. [Вложенная валидация](https://scikit-learn.org/stable/auto_examples/model_selection/plot_nested_cross_validation_iris.html), "
        "[Ridge Cox](https://scikit-survival.readthedocs.io/en/stable/user_guide/coxnet.html).", "",
        "C-index агрегирован по парам только внутри внешних блоков; IPCW C до трёх "
        "лет — среднее по внешним блокам с оценкой цензурирования только по их "
        "обучающей части. Исход и правила смерти совпадают с прежним протоколом.", ""]
    for scenario in SCENARIOS:
        info = first[scenario]
        lines += ["## " + ("Основная когорта" if scenario == "primary" else "Исключение death gap"), "",
            f"{info['patients']} пациентов / {info['distant_events']} событий. "
            "Среднее [минимум; максимум] по повторениям; разности вычислены парно "
            "на одинаковых внешних блоках.", "",
            "| Модель | C-index | IPCW C, 3 года | Разность C к Cox по объёму | Разность C к Ridge по объёму |",
            "| --- | ---: | ---: | ---: | ---: |"]
        for name in MODEL_COLUMNS:
            c, ipcw = (result["summary"][scenario][metric][name] for metric in METRICS)
            lines.append(f"| {LABELS[name]} | {fmt(c['scores'])} | {fmt(ipcw['scores'])} | "
                f"{fmt(c['paired_differences']['volume'])} | {fmt(c['paired_differences']['ridge_volume'])} |")
        lines += ["", "Выбранные alpha по внешним блокам (число выборов):", ""]
        for name in ("ridge_volume", "ridge_mri"):
            counts = Counter(fold["tuning"]["selected_alpha"] for repeat in result["repeats"]
                             for fold in repeat["scenarios"][scenario]["models"][name]["folds"])
            lines.append(f"- {LABELS[name]}: " + ", ".join(f"{a:g}: {n}" for a, n in sorted(counts.items())))
        lines.append("")
    lines += ["## Ограничения", "",
        "Панель добавлена после предыдущих экспериментов на той же когорте. Вложенная "
        "настройка защищает внешние блоки от подбора alpha, но не заменяет независимую "
        "проверку всей последовательности исследовательских решений. Повторения "
        "зависимы; их диапазоны не являются доверительными интервалами. "
        "Большая часть признаков не проверялась на воспроизводимость по исходным МРТ; "
        "доступность сегментации и различия сканеров остаются ограничениями. "
        "Сохраняются неопределённость исходов, возможное нарушение допущений "
        "цензурирования/пропорциональности рисков и малое число событий. Абсолютный "
        "риск, калибровка и внешняя валидация здесь не оценивались.", "",
        "Полные кандидаты alpha, внутренние метрики, статистики обработки, "
        "коэффициенты внешних моделей, версии и SHA-256: "
        "[mri_features_results.json](mri_features_results.json). Индивидуальные "
        "признаки, внешние/внутренние блоки и баллы сохраняются только в "
        "игнорируемой Git папке `data/derived/duke/mri_features/`.", ""]
    path.write_text("\n".join(lines))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline-dir", type=Path, default=Path("data/derived/duke/baseline"))
    parser.add_argument("--imaging-file", type=Path, default=Path("data/raw/duke/Imaging_Features.xlsx"))
    parser.add_argument("--data-dir", type=Path, default=Path("data/derived/duke/mri_features"))
    parser.add_argument("--output-dir", type=Path, default=Path("reports/duke"))
    args = parser.parse_args()
    started = time.perf_counter()
    baseline, outcomes, original_folds, manifest = load_prepared(args.baseline_dir)
    features = read_panel(args.imaging_file, baseline, manifest)
    rows = [{"patient_id": p, "event_type": int(outcomes.loc[p, "event_type"])} for p in features.index]
    if not np.array_equal(original_folds, pd.Series(assign_folds(rows)).loc[features.index]):
        raise ValueError("Original folds differ from the first planned outer split")
    args.data_dir.mkdir(parents=True, exist_ok=True)
    features.to_csv(args.data_dir / "features.csv")
    code_paths = [Path(__file__), *[Path(__file__).with_name(name) for name in (
        "run_duke_baseline.py", "prepare_duke_baseline.py", "check_duke_stability.py", "audit_duke.py")]]
    script_hashes = {p.name: sha256(p) for p in code_paths}
    note = Path("research/duke-mri-feature-rationale.md")
    inputs = {"baseline": manifest, "source_sha256": sha256(args.imaging_file),
        "panel": list(PANEL), "patients": len(features),
        "missing_by_feature": {c: int(features[c].isna().sum()) for c in PANEL},
        "rationale_sha256": sha256(note), "prepared_sha256": sha256(args.data_dir / "features.csv")}
    repeats, predictions, assignments, inner_rows = run_experiment(features, outcomes)
    if any(sha256(p) != script_hashes[p.name] for p in code_paths):
        raise ValueError("Experiment code changed during execution")
    result = {"status": "exploratory_nested_mri_panel_ranking_only", "inputs": inputs,
        "configuration": {"outer_seeds": list(SEEDS), "outer_n_folds": N_FOLDS,
            "inner_n_folds": INNER_FOLDS, "inner_seed_rule": "10 * outer_seed + outer_fold",
            "alphas": list(ALPHAS), "models": MODEL_COLUMNS, "horizon_days": HORIZON_DAYS,
            "selection_metric": "within-inner-fold pair-weighted Harrell C",
            "tie_break": "largest alpha within 1e-12 of best score",
            "feature_selection": "fixed nine-feature panel; no outcome-based filtering",
            "preprocessing": "median then standard scale fitted separately in each inner/outer training fold",
            "ridge_parameters_except_alpha": ridge_model(1.).steps[-1][1].get_params(),
            "baseline_parameters": make_model().steps[-1][1].get_params(),
            "death_handling": "same as baseline; death is censored, not a distant event",
            "sensitivity": "exclude death_gap, retain primary outer folds"},
        "software": {p: version(p) for p in ("scikit-survival", "scikit-learn", "numpy", "scipy", "pandas", "openpyxl")},
        "python": platform.python_version(), "script_sha256": script_hashes,
        "repeats": repeats, "summary": summarize(repeats), "elapsed_seconds": time.perf_counter() - started}
    save_csv(args.data_dir / "held_out_scores.csv", predictions,
        ["seed", "scenario", "model", "patient_id", "fold", "risk_score"])
    save_csv(args.data_dir / "outer_folds.csv", assignments, ["seed", "patient_id", "fold"])
    save_csv(args.data_dir / "inner_folds.csv", inner_rows,
        ["seed", "scenario", "outer_fold", "patient_id", "inner_fold"])
    result["artifact_sha256"] = {name: sha256(args.data_dir / name) for name in (
        "features.csv", "held_out_scores.csv", "outer_folds.csv", "inner_folds.csv")}
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "mri_features_inputs.json").write_text(json.dumps(inputs, indent=2, allow_nan=False) + "\n")
    (args.output_dir / "mri_features_results.json").write_text(json.dumps(result, indent=2, allow_nan=False) + "\n")
    write_report(args.output_dir / "mri_features.md", result)
    print(json.dumps(result["summary"]["primary"], indent=2))


if __name__ == "__main__":
    main()
